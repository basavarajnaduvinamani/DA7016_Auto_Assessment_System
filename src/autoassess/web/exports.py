"""Teacher downloads derived from current, audited marks."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile

from autoassess.classwork import class_report, class_report_html, consistency_pass
from autoassess.schemas import AnswerSegment, ExamSpec
from autoassess.web.grading import rebuild_artifacts


def grade_rows(store, assignment):
    spec = ExamSpec.model_validate_json(assignment['spec_json'])
    header = ['roll_no', 'name', *[q.id for q in spec.questions], 'total', 'percentage', 'reviewed_by_teacher']
    submissions = {s['student_id']: s for s in store.submissions(assignment['id']) if s['student_id']}
    rows = []
    for student in store.students(assignment['class_id']):
        submission = submissions.get(student['id'])
        grade = store.effective_grade(submission) if submission else None
        if grade:
            marks = {q.question_id: q.marks_awarded for q in grade.questions}
            reviewed = {q.id for q in spec.questions} <= store.accepted_questions(submission['id'])
            values = [marks.get(q.id, '') for q in spec.questions]
            values += [grade.total, round(grade.total / spec.total_marks * 100, 2), reviewed]
        else:
            values = [''] * (len(spec.questions) + 2) + [False]
        rows.append([student['roll_no'], student['name'], *values])
    return header, rows


def csv_bytes(header, rows):
    output = io.StringIO(newline='')
    writer = csv.writer(output)
    for row in [header, *rows]:
        # A name or imported question ID must not become an Excel formula.
        writer.writerow([('true' if cell else 'false') if isinstance(cell, bool) else
                         "'" + cell if isinstance(cell, str) and cell.lstrip().startswith(('=', '+', '-', '@')) else cell
                         for cell in row])
    return output.getvalue().encode('utf-8-sig')


def xlsx_bytes(header, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = 'Grade sheet'
    for row in [header, *rows]:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='AD343B')
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(42, max(14, max(len(str(c.value or '')) for c in column) + 2))
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def safe_pdf_name(roll_no, name):
    def clean(value):
        return re.sub(r'[/\\\x00-\x1f:*?"<>|]', '_', value).strip(' .')[:100] or 'student'
    return f'{clean(roll_no)}_{clean(name)}_marked.pdf'


def pdf_zip(store, assignment):
    output = io.BytesIO()
    names = set()
    with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for submission in store.submissions(assignment['id']):
            if not submission['grade_json']:
                continue
            artifacts = rebuild_artifacts(store, submission['id'])
            name = safe_pdf_name(submission['roll_no'], submission['name'])
            stem = name.removesuffix('_marked.pdf')
            suffix = 0
            while name in names:
                suffix += 1
                name = f"{stem}_{submission['id']}_{suffix}_marked.pdf"
            names.add(name)
            archive.write(artifacts['pdf'], name)
    return output.getvalue()


def report_content(store, assignment):
    spec = ExamSpec.model_validate_json(assignment['spec_json'])
    grades, segments = [], {}
    for submission in store.submissions(assignment['id']):
        grade = store.effective_grade(submission)
        if grade:
            grades.append(grade)
            saved = json.loads(submission['run_json'])
            segments[grade.student_id] = {qid: AnswerSegment.model_validate(segment) for qid, segment in saved['segments'].items()}
    checked = consistency_pass(spec, grades, segments)
    report = class_report(spec, checked)
    html = class_report_html(report, assignment['title'])
    body = html.split('<body>', 1)[1].rsplit('</body>', 1)[0]
    pairs = set()
    for grade in checked:
        for question in grade.questions:
            for flag in question.flags:
                if flag.startswith('inconsistent_with:'):
                    pair = tuple(sorted((grade.student_id, flag.split(':', 1)[1])))
                    pairs.add((question.question_id, *pair))
    return body, sorted(pairs)
