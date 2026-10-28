"""Current marks, spreadsheet safety, PDF downloads and printable class reports."""

import csv
import fnmatch
import io
import json
import socket
import zipfile

import pytest
import yaml
from conftest import ROOT
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from autoassess.schemas import ErrorMark, ExamSpec
from autoassess.web.app import create_app
from autoassess.web.exports import csv_bytes, safe_pdf_name, xlsx_bytes


@pytest.fixture
def export_app(tmp_path, monkeypatch):
    monkeypatch.setenv("AUTOASSESS_DEMO_DRY_RUN", "1")
    monkeypatch.setenv("GRADIO_ANALYTICS_ENABLED", "False")

    def blocked(*args, **kwargs):
        pytest.fail("Export tests must never call the network")

    monkeypatch.setattr(socket.socket, "connect", blocked)
    monkeypatch.setattr(socket, "getaddrinfo", blocked)
    app = create_app(tmp_path / "exports")
    with TestClient(app) as client:
        client.post("/sign-in", data={"email": "demo@redink.local"})
        yield app.state.store, client


def change_first(store):
    submission = store.submissions(1)[0]
    spec = ExamSpec.model_validate_json(store.assignment(1)["spec_json"])
    criterion = spec.questions[0].criteria[0]
    store.add_override(
        submission["id"], spec.questions[0].id, criterion.id, 0, "Teacher inspected the answer.", 1, expected_version=0
    )
    for q in spec.questions:
        store.accept_question(submission["id"], q.id, 1)
    return store.effective_grade(store.submission(submission["id"]))


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_grade_sheet_applies_override_and_preserves_ungraded_blanks(export_app, format):
    store, client = export_app
    expected = change_first(store)
    response = client.get(f"/assignments/1/export/{format}")
    assert response.status_code == 200 and "attachment" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    if format == "csv":
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    else:
        book = load_workbook(io.BytesIO(response.content))
        rows = list(book.active.values)
        assert book.active.freeze_panes == "C2"
    assert list(rows[0]) == ["roll_no", "name", "Q1", "Q2", "Q3", "total", "percentage", "reviewed_by_teacher"]
    assert len(rows) == 11  # every roster student, not only uploaded sheets
    first = rows[1]
    assert list(first[:2]) == ["S001", "Sample Student 01"]
    assert float(first[2]) == expected.questions[0].marks_awarded
    assert float(first[-3]) == expected.total
    assert float(first[-2]) == round(expected.total / 15 * 100, 2)
    assert first[-1] == (True if format == "xlsx" else "true")
    assert rows[4][0] == "S004" and all(cell in ("", None) for cell in rows[4][2:-1])
    assert rows[4][-1] == (False if format == "xlsx" else "false")
    assert rows[2][-1] == (False if format == "xlsx" else "false")


def test_spreadsheet_strings_cannot_execute_formulas():
    header = ["=header", "name", "total"]
    rows = [["S001", '=HYPERLINK("http://invalid")', 1.5], ["S002", "  +SUM(1,2)", 0]]
    csv_rows = list(csv.reader(io.StringIO(csv_bytes(header, rows).decode("utf-8-sig"))))
    assert csv_rows[0][0] == "'=header"
    assert csv_rows[1][1].startswith("'=") and csv_rows[2][1].startswith("'  +")
    book = load_workbook(io.BytesIO(xlsx_bytes(header, rows)))
    assert book.active["A1"].data_type == "s" and book.active["B2"].data_type == "s"
    assert book.active["B2"].value == rows[0][1] and book.active["C2"].value == 1.5


def test_pdf_zip_names_and_current_artifacts(export_app, monkeypatch):
    from autoassess.web import exports

    store, client = export_app
    changed = change_first(store)
    calls = []
    original = exports.rebuild_artifacts

    def rebuild(database, sid):
        calls.append((sid, database.effective_grade(database.submission(sid)).total))
        return original(database, sid)

    monkeypatch.setattr(exports, "rebuild_artifacts", rebuild)
    response = client.get("/assignments/1/export/zip")
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert archive.namelist() == [f"S00{i}_Sample Student 0{i}_marked.pdf" for i in (1, 2, 3)]
        assert all(archive.read(name).startswith(b"%PDF") for name in archive.namelist())
    assert calls[0][1] == changed.total and len(calls) == 3
    assert "/" not in safe_pdf_name("S001", "../../Bad:Name") and "\\" not in safe_pdf_name("S001", "Bad\\Name")


def test_class_report_reuses_current_marks_and_shows_consistency_pairs(export_app):
    store, client = export_app
    spec = ExamSpec.model_validate_json(store.assignment(1)["spec_json"])
    submissions = store.submissions(1)
    for submission, marks in zip(submissions[:2], [True, False], strict=True):
        saved = json.loads(submission["run_json"])
        saved["segments"]["Q1"]["text"] = "Same synthetic answer, deliberately assigned different marks."
        with store.connect() as db:
            db.execute("UPDATE submissions SET run_json=? WHERE id=?", (json.dumps(saved), submission["id"]))
        for criterion in spec.questions[0].criteria:
            changes = store.changes(submission["id"])
            store.add_override(
                submission["id"],
                "Q1",
                criterion.id,
                criterion.marks if marks else 0,
                "Consistency example",
                1,
                expected_version=changes[-1]["id"] if changes else 0,
            )
    grade = store.effective_grade(store.submission(submissions[2]["id"]))
    grade.questions[0].errors = [ErrorMark(quote="<unsafe>", correction="Check the definition")]
    with store.connect() as db:
        db.execute("UPDATE submissions SET grade_json=? WHERE id=?", (grade.model_dump_json(), submissions[2]["id"]))
    response = client.get("/assignments/1/report")
    assert response.status_code == 200 and "Class report" in response.text
    assert "Common mistakes" in response.text and "Check the definition" in response.text
    assert "&lt;unsafe&gt;" in response.text and "<unsafe>" not in response.text
    assert "Similar answers with different marks" in response.text
    assert "<td>S001</td><td>S002</td>" in response.text
    assert "Print class report" in response.text
    assert "@media print" in (ROOT / "src/autoassess/web/static/app.css").read_text()


def test_download_ownership_and_unknown_format(export_app):
    _, client = export_app
    assert client.get("/assignments/999/export/csv").status_code == 404
    assert client.get("/assignments/1/export/exe").status_code == 404
    client.post("/sign-out")
    for path in [
        "/assignments/1/export/csv",
        "/assignments/1/export/xlsx",
        "/assignments/1/export/zip",
        "/assignments/1/report",
    ]:
        assert client.get(path, follow_redirects=False).status_code == 303


def test_docker_context_excludes_private_files_and_builds_in_ci():
    patterns = [
        line.strip()
        for line in (ROOT / ".dockerignore").read_text().splitlines()
        if line.strip() and not line.startswith("#")
    ]

    def ignored(path):
        pieces = path.split("/")
        return any(
            fnmatch.fnmatch("/".join(pieces[:i]), pattern) for i in range(1, len(pieces) + 1) for pattern in patterns
        )

    for path in [
        "data/private/real.csv",
        "data/private/scans/student.pdf",
        "runs/run/grade.json",
        ".cache/ledger.json",
        ".redink/redink.db",
        ".redink/files/sheet.pdf",
        ".env",
        ".env.production",
        ".git/config",
        "nested/data/private/answer.txt",
    ]:
        assert ignored(path), path
    for path in [
        "pyproject.toml",
        "README.md",
        "src/autoassess/web/app.py",
        "src/autoassess/sample_exam/exam_spec.yaml",
    ]:
        assert not ignored(path), path
    dockerfile = (ROOT / "Dockerfile").read_text()
    assert "FROM python:3.12-slim" in dockerfile and ".[web,ingest,annotate]" in dockerfile
    assert "USER redink" in dockerfile and "autoassess serve --host 0.0.0.0 --port" in dockerfile
    workflow = yaml.load((ROOT / ".github/workflows/ci.yml").read_text(), Loader=yaml.BaseLoader)
    steps = workflow["jobs"]["docker"]["steps"]
    assert any("docker build" in step.get("run", "") for step in steps)
    assert not any("docker push" in step.get("run", "") for step in steps)
    assert any("--network none" in step.get("run", "") for step in steps)
    render = yaml.safe_load((ROOT / "render.yaml").read_text())["services"][0]
    assert render["startCommand"] == "autoassess serve --host 0.0.0.0 --port $PORT"
    assert render["healthCheckPath"] == "/healthz"


def test_zip_name_collisions_keep_every_pdf(tmp_path, monkeypatch):
    from types import SimpleNamespace

    import autoassess.web.exports as exports

    pdf = tmp_path / "marked.pdf"
    pdf.write_bytes(b"%PDF-synthetic test artifact")
    rows = [
        {"id": 1, "student_id": 1, "roll_no": "A_B", "name": "C", "grade_json": "{}"},
        {"id": 2, "student_id": 2, "roll_no": "A", "name": "B_C", "grade_json": "{}"},
    ]
    database = SimpleNamespace(submissions=lambda aid: rows)
    monkeypatch.setattr(exports, "rebuild_artifacts", lambda store, sid: {"pdf": str(pdf)})
    with zipfile.ZipFile(io.BytesIO(exports.pdf_zip(database, {"id": 1}))) as archive:
        assert len(set(archive.namelist())) == len(rows)
        assert all(name.endswith("_marked.pdf") and "/" not in name for name in archive.namelist())
